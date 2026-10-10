from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
import struct, zlib
p = Presentation()
p.slide_width = Emu(12192000); p.slide_height = Emu(6858000)
s = p.slides.add_slide(p.slide_layouts[0])
s.shapes.title.text = '네이버 검색광고 성장 전략 제안서'
s.placeholders[1].text = '성과 진단, 캠페인 재구조화\n그리고 압도적 성장을 위한 로드맵'
s2 = p.slides.add_slide(p.slide_layouts[1])
s2.shapes.title.text = '코어 키워드 카테고리 (Core Keywords)'
tf = s2.placeholders[1].text_frame
tf.text = '스마트 주차 솔루션'
for t in ['주차 관제 시스템', 'Bold English Words', '무인 주차장 운영']:
    para = tf.add_paragraph(); para.text = t
r = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8), Inches(2), Inches(3), Inches(1.5))
r.fill.solid(); r.fill.fore_color.rgb = RGBColor(0x2E, 0x4B, 0xA0)
r.text_frame.text = '70 Core'
o = s2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8), Inches(4), Inches(2), Inches(2))
o.fill.solid(); o.fill.fore_color.rgb = RGBColor(0xF0, 0x80, 0x20)
# png image
w,h=64,32
raw=b''.join(b'\x00'+b''.join(bytes([x*4%256,y*8%256,128]) for x in range(w)) for y in range(h))
def chunk(t,d): return struct.pack('>I',len(d))+t+d+struct.pack('>I',zlib.crc32(t+d)&0xffffffff)
png=b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR',struct.pack('>IIBBBBB',w,h,8,2,0,0,0))+chunk(b'IDAT',zlib.compress(raw))+chunk(b'IEND',b'')
open('img.png','wb').write(png)
s3 = p.slides.add_slide(p.slide_layouts[5])
s3.shapes.title.text = '그림과 차트'
s3.shapes.add_picture('img.png', Inches(1), Inches(2), Inches(4), Inches(2))
cd = CategoryChartData(); cd.categories=['1월','2월','3월']; cd.add_series('클릭', (10, 25, 18))
s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6), Inches(2), Inches(6), Inches(4), cd)
p.save('t.pptx')
