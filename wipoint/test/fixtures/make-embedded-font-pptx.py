import zipfile, re
from pptx import Presentation
p = Presentation()
s = p.slides.add_slide(p.slide_layouts[5])
s.shapes.title.text = 'Embedded Font Test 123'
for r in s.shapes.title.text_frame.paragraphs[0].runs: r.font.name = 'WP Embedded Serif'
p.save('/tmp/_base.pptx')
zin = zipfile.ZipFile('/tmp/_base.pptx')
zout = zipfile.ZipFile('embedded-font.pptx', 'w', zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    d = zin.read(it.filename)
    if it.filename == '[Content_Types].xml':
        d = d.replace(b'<Default Extension="xml"', b'<Default Extension="fntdata" ContentType="application/x-fontdata"/><Default Extension="xml"')
    if it.filename == 'ppt/_rels/presentation.xml.rels':
        d = d.replace(b'</Relationships>', b'<Relationship Id="rIdF1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/font" Target="fonts/font1.fntdata"/></Relationships>')
    if it.filename == 'ppt/presentation.xml':
        d = re.sub(rb'<p:presentation ', b'<p:presentation embedTrueTypeFonts="1" ', d, count=1)
        d = re.sub(rb'(<p:notesSz[^>]*/>)', rb'\1<p:embeddedFontLst><p:embeddedFont><p:font typeface="WP Embedded Serif" pitchFamily="18" charset="0"/><p:regular r:id="rIdF1"/></p:embeddedFont></p:embeddedFontLst>', d, count=1)
    zout.writestr(it, d)
zout.writestr('ppt/fonts/font1.fntdata', open('sub.eot', 'rb').read())
zout.close()
