from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
p = Presentation()
s = p.slides.add_slide(p.slide_layouts[0])
s.shapes.title.text = 'python-pptx 제목'
s.placeholders[1].text = '부제목 텍스트'
s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = '글머리 기호'
tf = s.placeholders[1].text_frame
tf.text = '첫째 항목'
for t, l in [('둘째 수준', 1), ('셋째 항목', 0)]:
    para = tf.add_paragraph(); para.text = t; para.level = l
s.notes_slide.notes_text_frame.text = '발표 메모 python'
s = p.slides.add_slide(p.slide_layouts[5])
s.shapes.title.text = '도형 · 표 · 차트'
sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.5), Inches(2.5), Inches(1.2))
sh.text = '기본 스타일 도형'
sh2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(3), Inches(1.5), Inches(1.5))
sh2.fill.solid(); sh2.fill.fore_color.rgb = RGBColor(0xC0, 0x00, 0x00)
tb = s.shapes.add_table(3, 3, Inches(3.3), Inches(1.5), Inches(3), Inches(1.2)).table
tb.cell(0, 0).text = '헤더'; tb.cell(1, 1).text = '42'
cd = CategoryChartData(); cd.categories = ['가', '나', '다']; cd.add_series('매출', (3, 5, 2))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.5), Inches(1.5), Inches(3.2), Inches(2.5), cd)
g = s.shapes.add_group_shape()
g.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(3.5), Inches(3.2), Inches(1), Inches(1))
g.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.6), Inches(3.4), Inches(1.2), Inches(0.6))
p.save('pp.pptx')
print('ok')
